import argparse
import csv
import hashlib
import json
import sys
from pathlib import Path

from .core import digest, match, read_json, tailor, validate_job, validate_profile
from .export import export_bundle
from .store import Store


def emit(value):
    print(json.dumps(value, ensure_ascii=False, indent=2))


def main():
    p = argparse.ArgumentParser(description="JobPilot 本地求职工作台：真实经历 → 岗位匹配 → 定制简历 → 可追踪投递")
    p.add_argument("--db", default="data/jobpilot.sqlite")
    sub = p.add_subparsers(dest="cmd", required=True)
    cmd = sub.add_parser("extract", help="提取 TXT/PDF/DOCX 原文供整理，不自动认证经历")
    cmd.add_argument("file")
    cmd.add_argument("--out", required=True)
    cmd = sub.add_parser("import-jobs", help="导入岗位 JSON 数组或单个对象")
    cmd.add_argument("file")
    cmd = sub.add_parser("prepare", help="筛选并生成逐岗位投递包")
    cmd.add_argument("--profile", required=True)
    cmd.add_argument("--policy", required=True)
    cmd.add_argument("--out", default="output")
    cmd.add_argument("--ai", action="store_true", help="把经历事实与岗位要求发至自己配置的模型接口")
    cmd.add_argument("--docx", action="store_true")
    cmd = sub.add_parser("status", help="查看投递状态")
    cmd.add_argument("--csv")
    cmd = sub.add_parser("events", help="查看某岗位完整事件记录")
    cmd.add_argument("job_id")
    cmd = sub.add_parser("login", help="打开本机专用浏览器，手动登录")
    cmd.add_argument("url")
    cmd.add_argument("--session", default="browser-profile")
    cmd = sub.add_parser("collect", help="按已配置页面定位器采集岗位")
    cmd.add_argument("url")
    cmd.add_argument("--adapter", required=True)
    cmd.add_argument("--session", default="browser-profile")
    cmd.add_argument("--pages", type=int, default=1)
    cmd.add_argument("--out", required=True)
    cmd = sub.add_parser("run", help="默认只核对页面；--execute 执行本批次投递")
    cmd.add_argument("--adapter", required=True)
    cmd.add_argument("--session", default="browser-profile")
    cmd.add_argument("--limit", type=int, default=5)
    cmd.add_argument("--daily-limit", type=int, default=10)
    cmd.add_argument("--execute", action="store_true")
    a = p.parse_args()
    db = Store(a.db)
    try:
        if a.cmd == "extract":
            path = Path(a.file)
            if path.suffix.lower() == ".pdf":
                from pypdf import PdfReader
                text = "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
            elif path.suffix.lower() == ".docx":
                from docx import Document
                d = Document(path)
                text = "\n".join([p.text for p in d.paragraphs] + [" | ".join(c.text for c in row.cells) for t in d.tables for row in t.rows])
            elif path.suffix.lower() in (".txt", ".md"):
                text = path.read_text(encoding="utf-8-sig")
            else:
                raise ValueError("仅支持 TXT、MD、PDF、DOCX")
            if not text.strip():
                raise ValueError("未提取出文字，扫描 PDF 需要先 OCR")
            Path(a.out).parent.mkdir(parents=True, exist_ok=True)
            Path(a.out).write_text(text, encoding="utf-8")
            emit({"extracted": a.out, "next": "按 examples/profile.json 整理并核对真实经历"})
        elif a.cmd == "import-jobs":
            jobs = read_json(a.file)
            jobs = jobs if isinstance(jobs, list) else [jobs]
            jobs = [validate_job(j) for j in jobs]  # validate batch before first write
            for j in jobs:
                db.put_job(j)
            emit({"imported": len(jobs), "unique_total": len(db.jobs())})
        elif a.cmd == "prepare":
            profile = validate_profile(read_json(a.profile))
            policy = read_json(a.policy)
            results = []
            for job in db.jobs():
                old = db.application(job["id"])
                if old and old["status"] not in ("matched", "needs_review", "rejected"):
                    results.append({"job_id": job["id"], "skipped": old["status"]})
                    continue
                assessment = match(profile, job, policy)
                selection = None
                if a.ai and assessment["status"] == "matched":
                    from .llm import select_facts
                    selection = select_facts(profile, job, assessment)
                bundle = tailor(profile, job, assessment, selection)
                bundle["policy_hash"] = digest(policy)
                directory = Path(a.out).resolve() / job["id"] / digest(bundle)[:16]
                # Snapshot directories retain older resume versions for audit.
                export_bundle(bundle, directory, a.docx)
                bundle["attachment_hashes"] = {x.name: hashlib.sha256(x.read_bytes()).hexdigest() for x in directory.glob("resume.*") if x.suffix in (".docx", ".pdf")}
                path = directory / "bundle.json"
                path.write_text(json.dumps(bundle, ensure_ascii=False, indent=2), encoding="utf-8")
                db.save_bundle(job["id"], path, digest(bundle), assessment["status"])
                results.append({"job_id": job["id"], "title": job["title"], **assessment, "resume": str(directory / "resume.html")})
            emit(results)
        elif a.cmd == "status":
            rows = db.applications()
            if a.csv:
                Path(a.csv).parent.mkdir(parents=True, exist_ok=True)
                with open(a.csv, "w", encoding="utf-8-sig", newline="") as f:
                    w = csv.DictWriter(f, fieldnames=["job_id", "status", "bundle_path", "bundle_hash", "updated_at"])
                    w.writeheader()
                    w.writerows(rows)
            emit(rows)
        elif a.cmd == "events":
            emit([dict(r) for r in db.db.execute("SELECT * FROM events WHERE job_id=? ORDER BY id", (a.job_id,))])
        elif a.cmd == "login":
            from .browser import login
            login(a.url, a.session)
        elif a.cmd == "collect":
            from .browser import collect
            jobs = collect(a.url, read_json(a.adapter), a.session, a.pages)
            Path(a.out).parent.mkdir(parents=True, exist_ok=True)
            Path(a.out).write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")
            for job in jobs:
                db.put_job(job)
            emit({"collected": len(jobs), "file": a.out})
        elif a.cmd == "run":
            from .browser import execute_queue
            emit(execute_queue(db, read_json(a.adapter), a.session, a.limit, a.daily_limit, a.execute))
    except Exception as e:
        # Do not echo arbitrary HTTP bodies, tokens, or candidate data into logs.
        if isinstance(e, (ValueError, FileNotFoundError, ImportError)):
            print(f"错误：{e}", file=sys.stderr)
        else:
            print(f"操作中止：{type(e).__name__}。请核对配置、依赖或页面；未自动重试。", file=sys.stderr)
        return_code = 1
    else:
        return_code = 0
    finally:
        db.close()
    if return_code:
        raise SystemExit(return_code)
