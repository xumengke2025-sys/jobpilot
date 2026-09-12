import html
import json
from pathlib import Path

STAGES = {"production": "已落地", "pilot": "试点", "prototype": "原型", "research": "研究规划"}


def render_resume(bundle):
    esc = lambda v: html.escape(str(v))
    # Preserve the supplied employment timeline. Relevant projects may be re-ordered.
    history = "".join(f"<p>{esc(x)}</p>" for x in bundle.get("employment", []))
    education = "".join(f"<p>{esc(x)}</p>" for x in bundle.get("education", []))
    facts = "".join(f"<article><h3>{esc(f['project'])} <small>{STAGES[f['stage']]}</small></h3><p class='meta'>{esc(f['organization'])} · {esc(f['role'])} · {esc(f['period'])}</p><p>{esc(f['text'])}</p></article>" for f in bundle["facts"])
    return f"""<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{esc(bundle['name'])} · 简历</title>
<style>body{{font:15px/1.75 'Microsoft YaHei',sans-serif;color:#172b39;max-width:820px;margin:40px auto;padding:0 28px}}h1{{margin-bottom:0;font-size:30px}}h2{{font-size:18px;border-bottom:1px solid #b7c8cc;padding-bottom:6px;margin-top:26px}}h3{{font-size:16px;margin-bottom:3px}}small,.meta{{font-size:12px;color:#526971}}article{{break-inside:avoid}}.skills{{color:#176b65}}@media print{{body{{margin:0;padding:0;font-size:11pt}}@page{{size:A4;margin:18mm}}}}</style>
<h1>{esc(bundle['name'])}</h1><p>{esc(bundle['headline'])}<br>{esc(bundle.get('contact',''))}</p>
<h2>相关能力</h2><p class="skills">{esc(' · '.join(bundle['keywords']))}</p>
<h2>工作经历</h2>{history}<h2>相关项目</h2>{facts}<h2>教育经历</h2>{education}</html>"""


def export_bundle(bundle, directory, docx=False):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    (directory / "resume.html").write_text(render_resume(bundle), encoding="utf-8")
    (directory / "bundle.json").write_text(json.dumps(bundle, ensure_ascii=False, indent=2), encoding="utf-8")
    (directory / "greeting.txt").write_text(bundle["greeting"], encoding="utf-8")
    (directory / "match.json").write_text(json.dumps(bundle["assessment"], ensure_ascii=False, indent=2), encoding="utf-8")
    if docx:
        from docx import Document
        from docx.shared import Pt
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        d = Document()
        style = d.styles["Normal"]
        style.font.name, style.font.size = "Microsoft YaHei", Pt(10.5)
        rpr = style.element.get_or_add_rPr()
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        rpr.append(fonts)
        d.add_heading(bundle["name"], 0)
        d.add_paragraph(bundle["headline"])
        d.add_paragraph(bundle.get("contact", ""))
        d.add_heading("相关能力", 1)
        d.add_paragraph(" · ".join(bundle["keywords"]))
        d.add_heading("工作经历", 1)
        for line in bundle.get("employment", []):
            d.add_paragraph(str(line))
        d.add_heading("相关项目", 1)
        for f in bundle["facts"]:
            d.add_heading(f["project"] + " · " + STAGES[f["stage"]], 2)
            d.add_paragraph(f"{f['organization']} · {f['role']} · {f['period']}")
            d.add_paragraph(f["text"])
        d.add_heading("教育经历", 1)
        for line in bundle.get("education", []):
            d.add_paragraph(str(line))
        d.save(directory / "resume.docx")
    return directory / "bundle.json"
